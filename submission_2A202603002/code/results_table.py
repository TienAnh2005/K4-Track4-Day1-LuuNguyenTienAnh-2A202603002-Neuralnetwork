"""results_table.py — Lưu kết quả và xuất bảng Excel experiments.xlsx.

Chức năng:
  - save_result: lưu cfg, history, summary thành file json
  - load_results: đọc danh sách json trong thư mục results
  - to_row: chuyển kết quả thành 1 dòng cho sheet Experiments
  - write_xlsx: ghi vào template và cập nhật công thức
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary ra file JSON."""
    exp_id = result.get("cfg", {}).get("exp_id", "exp")
    p = Path(results_dir)
    p.mkdir(parents=True, exist_ok=True)
    out_file = p / f"{exp_id}.json"

    data_to_save = {
        "cfg": result.get("cfg", {}),
        "history": result.get("history", {}),
        "summary": result.get("summary", {}),
    }

    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(out_file)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict."""
    p = Path(results_dir)
    if not p.exists():
        return []

    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            data = json.load(fp)
            results.append(data)

    results.sort(key=lambda r: r.get("cfg", {}).get("exp_id", ""))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của sheet Experiments."""
    cfg = result.get("cfg", {})
    summary = result.get("summary", {})
    exp_id = cfg.get("exp_id", "")

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr", ""),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": str(tuple(cfg.get("hidden", (256, 128)))),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", ""),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", ""),
        "best_val_loss": summary.get("best_val_loss", ""),
        "best_epoch": summary.get("best_epoch", ""),
        "final_train_loss": summary.get("final_train_loss", ""),
        "final_val_loss": summary.get("final_val_loss", ""),
        "val_acc": summary.get("val_acc", ""),
        "val_macro_f1": summary.get("val_macro_f1", ""),
        "time_per_epoch_s": summary.get("time_per_epoch_s", ""),
        "peak_mem_MB": summary.get("peak_mem_MB", ""),
        "diverged": summary.get("diverged", False),
        "eval_acc": eval_scores.get("acc", "") if eval_scores else "",
        "eval_macro_f1": eval_scores.get("macro_f1", "") if eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet Experiments của template và lưu ra out_path."""
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    header_cols = {}
    for col_idx in range(1, ws.max_column + 1):
        name = ws.cell(row=1, column=col_idx).value
        if name:
            header_cols[str(name).strip()] = col_idx

    formula_cols = {
        "step0_gap_vs_lnC": lambda r: f'=IF(P{r}="","",P{r}-LN(7))',
        "gap_val_minus_train": lambda r: f'=IF(OR(T{r}="",S{r}=""),"",T{r}-S{r})',
        "delta_val_f1_vs_base": lambda r: f'=IF(OR(V{r}="",Seeds!$C$8=""),"",V{r}-Seeds!$C$8)',
        "beyond_noise": lambda r: f'=IF(OR(AF{r}="",Seeds!$C$10=""),"",IF(ABS(AF{r})>Seeds!$C$10,"Có","Không"))',
    }

    start_row = 2
    for i, row_data in enumerate(rows):
        curr_row = start_row + i
        for key, val in row_data.items():
            if key in header_cols:
                ws.cell(row=curr_row, column=header_cols[key]).value = val

        for form_key, form_fn in formula_cols.items():
            if form_key in header_cols:
                ws.cell(row=curr_row, column=header_cols[form_key]).value = form_fn(curr_row)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Đã ghi {len(rows)} thí nghiệm vào {out_path}")
